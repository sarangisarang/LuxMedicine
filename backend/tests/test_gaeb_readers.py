"""Readers turn a source file into positions. The subtle part is numbers (German vs English) and
finding the header row; the strongest guarantee is the GAEB round-trip — what the writer emits, the
reader reads back unchanged."""

import io
from datetime import datetime
from decimal import Decimal

import openpyxl
import pytest

from app.services.gaeb.model import BillOfQuantities, Position
from app.services.gaeb.readers import (
    NoPositionsError,
    UnreadableFileError,
    parse_decimal,
    read_csv,
    read_docx,
    read_gaeb,
    read_pdf,
    read_xlsx,
)
from app.services.gaeb.writer import write_x84


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("1.234,56", Decimal("1234.56")),   # German
        ("1,234.56", Decimal("1234.56")),   # English
        ("1234,56", Decimal("1234.56")),    # comma decimal, no thousands
        ("1234.56", Decimal("1234.56")),    # plain
        ("€ 125,50", Decimal("125.50")),    # currency symbol
        ("10", Decimal("10")),
        ("", None),
        (None, None),
    ],
)
def test_parse_decimal_reads_both_conventions(raw, expected):
    assert parse_decimal(raw) == expected


def test_parse_decimal_rejects_nonsense():
    with pytest.raises(UnreadableFileError):
        parse_decimal("not a number")


CSV_DE = (
    "Pos;Kurztext;Menge;Einheit;Einheitspreis\n"
    "01.0010;Beton C25/30;10;m3;125,50\n"
    "01.0020;Bewehrungsstahl;2,5;t;900,00\n"
)


def test_read_csv_german_semicolons():
    boq = read_csv(CSV_DE.encode("utf-8"), project_name="Halle")
    assert boq.project_name == "Halle"
    assert len(boq.positions) == 2
    p = boq.positions[0]
    assert (p.oz, p.short_text, p.unit) == ("01.0010", "Beton C25/30", "m3")
    assert p.quantity == Decimal("10")
    assert p.unit_price == Decimal("125.50")
    assert boq.positions[1].quantity == Decimal("2.5")


def test_read_csv_empty_is_unreadable():
    with pytest.raises(UnreadableFileError):
        read_csv(b"   \n  \n", project_name="x")


def test_read_csv_unknown_columns_has_no_positions():
    with pytest.raises(NoPositionsError):
        read_csv(b"Alpha;Beta;Gamma\n1;2;3\n", project_name="x")


def _xlsx(rows: list[list[object]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_xlsx_skips_title_rows_above_the_header():
    data = _xlsx([
        ["Leistungsverzeichnis Neubau"],          # a title row, no table
        [],                                        # a blank row
        ["OZ", "Bezeichnung", "Menge", "ME", "EP"],
        ["1.1", "Mauerwerk", 12.5, "m2", 45.0],
        ["1.2", "Putz", 30, "m2", 18.5],
    ])
    boq = read_xlsx(data, project_name="Neubau")
    assert len(boq.positions) == 2
    assert boq.positions[0].short_text == "Mauerwerk"
    assert boq.positions[0].quantity == Decimal("12.5")
    assert boq.positions[1].unit_price == Decimal("18.5")


def test_read_xlsx_without_a_recognisable_header_raises():
    data = _xlsx([["foo", "bar"], [1, 2]])
    with pytest.raises(NoPositionsError):
        read_xlsx(data, project_name="x")


def test_gaeb_round_trip_preserves_positions():
    boq = BillOfQuantities(
        project_name="Rathaus",
        positions=(
            Position(oz="01.0010", short_text="Beton", quantity=Decimal("10"), unit="m3",
                     unit_price=Decimal("125.50"), section="Rohbau"),
            Position(oz="01.0020", short_text="Stahl", quantity=Decimal("2.5"), unit="t",
                     unit_price=Decimal("900.00"), section="Rohbau"),
        ),
    )
    xml = write_x84(boq, generated_at=datetime(2026, 7, 25, 12, 0, 0))

    back = read_gaeb(xml)
    assert back.project_name == "Rathaus"
    assert len(back.positions) == 2
    a = back.positions[0]
    assert a.oz == "01.0010"
    assert a.short_text == "Beton"
    assert a.quantity == Decimal("10")
    assert a.unit == "m3"
    assert a.unit_price == Decimal("125.50")
    assert a.section == "Rohbau"


def test_read_gaeb_rejects_non_gaeb_xml():
    with pytest.raises(UnreadableFileError):
        read_gaeb(b"<html><body>not gaeb</body></html>")


def test_unrecognised_headers_fall_back_to_reading_the_data():
    """The real-world failure: an LV whose columns are labelled in a way the list does not know.
    The data still says which column is which — units are units, numbers are numbers."""
    csv = (
        "Kennung;Was zu tun ist;Umf.;Einh;Satz\n"
        "01.0010;Beton C25/30 liefern und einbauen;10;m3;125,50\n"
        "01.0020;Bewehrungsstahl verlegen;2,5;t;900,00\n"
    ).encode("utf-8")
    boq = read_csv(csv, project_name="Halle")
    assert len(boq.positions) == 2
    assert boq.positions[0].short_text.startswith("Beton")
    assert boq.positions[0].quantity == Decimal("10")
    assert boq.positions[0].unit == "m3"
    assert boq.positions[0].unit_price == Decimal("125.50")


def test_unrecognised_headers_with_only_one_position():
    """The case that failed in production while the two-row test above passed: with a single
    position, the unnamed header row's own text drags every column below the numeric threshold. The
    inference has to try again without that first row."""
    csv = (
        "Kennung;Was zu tun ist;Umf.;Einh;Satz\n"
        "01.0010;Beton liefern und einbauen;10;m3;125,50\n"
    ).encode("utf-8")
    boq = read_csv(csv, project_name="Halle")
    assert len(boq.positions) == 1
    assert boq.positions[0].quantity == Decimal("10")
    assert boq.positions[0].unit == "m3"
    assert boq.positions[0].unit_price == Decimal("125.50")


def test_a_table_with_no_header_row_at_all_is_read_from_its_content():
    """Word and PDF LVs frequently print no header. Every row is a position; the columns are read
    from what they contain."""
    csv = (
        "01.0010;Beton C25/30 liefern und einbauen;10;m3;125,50\n"
        "01.0020;Bewehrungsstahl verlegen;2,5;t;900,00\n"
    ).encode("utf-8")
    boq = read_csv(csv, project_name="Halle")
    assert len(boq.positions) == 2
    assert boq.positions[1].unit == "t"
    assert boq.positions[1].quantity == Decimal("2.5")


def test_a_missing_unit_column_is_tolerated():
    """A .x84 position can carry an empty unit; a description and a quantity are the real minimum."""
    csv = b"Pos;Bezeichnung;Menge;EP\n1;Malerarbeiten;5;20,00\n"
    boq = read_csv(csv, project_name="x")
    assert len(boq.positions) == 1
    assert boq.positions[0].quantity == Decimal("5")


def test_failure_says_what_it_actually_read():
    """A blind 'could not find a column' teaches nothing. The message must echo the rows it saw."""
    with pytest.raises(NoPositionsError) as excinfo:
        read_csv(b"Alpha;Beta\nnur Text;auch Text\n", project_name="x")
    message = str(excinfo.value)
    assert "Alpha" in message or "nur Text" in message


def _docx(header: list[str], rows: list[list[str]]) -> bytes:
    import docx

    document = docx.Document()
    table = document.add_table(rows=1 + len(rows), cols=len(header))
    for j, h in enumerate(header):
        table.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            table.rows[i].cells[j].text = str(value)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_read_docx_table():
    data = _docx(
        ["Pos", "Bezeichnung", "Menge", "ME", "EP"],
        [["1.1", "Mauerwerk", "12,5", "m2", "45,00"], ["1.2", "Putz", "30", "m2", "18,50"]],
    )
    boq = read_docx(data, project_name="Neubau")
    assert len(boq.positions) == 2
    assert boq.positions[0].short_text == "Mauerwerk"
    assert boq.positions[0].quantity == Decimal("12.5")
    assert boq.positions[1].unit_price == Decimal("18.5")


def test_read_docx_without_a_table_raises():
    import docx

    document = docx.Document()
    document.add_paragraph("Just prose, no table.")
    buf = io.BytesIO()
    document.save(buf)
    with pytest.raises(NoPositionsError):
        read_docx(buf.getvalue(), project_name="x")


def _pdf_with_table(header: list[str], rows: list[list[str]]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    table = Table([header, *rows])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc.build([table])
    return buf.getvalue()


def test_read_pdf_ruled_table():
    data = _pdf_with_table(
        ["Pos", "Bezeichnung", "Menge", "ME", "EP"],
        [["1.1", "Mauerwerk", "12,5", "m2", "45,00"], ["1.2", "Putz", "30", "m2", "18,50"]],
    )
    boq = read_pdf(data, project_name="Neubau")
    assert len(boq.positions) == 2
    assert boq.positions[0].short_text == "Mauerwerk"
    assert boq.positions[0].quantity == Decimal("12.5")


def test_read_pdf_without_a_table_reports_cleanly():
    # A PDF with only prose (no ruled table) — the honest "couldn't extract", not a garbled parse.
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    from reportlab.lib.styles import getSampleStyleSheet

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4).build(
        [Paragraph("A letter with no bill of quantities.", getSampleStyleSheet()["Normal"])]
    )
    with pytest.raises(NoPositionsError):
        read_pdf(buf.getvalue(), project_name="x")
